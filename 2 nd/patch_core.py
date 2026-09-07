import os

target = r"C:\Users\habee\backend\app\api\core.py"

with open(target, "r", encoding="utf-8") as f:
    text = f.read()

# 1. Update ALIASES to include date_of_record, district, state, station_name, city, air_pressure, etc.
old_aliases = "ALIASES = {'date':['date','time','datetime'], 'region':['region','area','location'], 'rainfall':['rainfall','rain','precipitation','precip_mm','rainfall_mm'], 'temperature':['temperature','temp'], 'humidity':['humidity'], 'wind_speed':['wind_speed','windspeed','wind'], 'pressure':['pressure','atmospheric_pressure'], 'cloud_cover':['cloud_cover','cloud'], 'water_consumption':['water_consumption','water_usage','water_demand','consumption'], 'population':['population'], 'reservoir_level':['reservoir_level','reservoir']}"

new_aliases = "ALIASES = {'date':['date','time','datetime','date_of_record'], 'region':['region','area','location','district','state','station_name','city'], 'rainfall':['rainfall','rain','precipitation','precip_mm','rainfall_mm'], 'temperature':['temperature','temp','avg_temp'], 'humidity':['humidity'], 'wind_speed':['wind_speed','windspeed','wind'], 'pressure':['pressure','atmospheric_pressure','air_pressure'], 'cloud_cover':['cloud_cover','cloud'], 'water_consumption':['water_consumption','water_usage','water_demand','consumption'], 'population':['population'], 'reservoir_level':['reservoir_level','reservoir']}"

text = text.replace(old_aliases, new_aliases)

# 2. Update load and metadata functions
old_load_meta = '''def load(path):
    df = pd.read_excel(path) if path.lower().endswith(('.xlsx','.xls')) else pd.read_csv(path)
    mapped={}
    for col in df.columns:
        n=norm(col)
        for target, names in ALIASES.items():
            if n in names or any(x in n for x in names): mapped[col]=target; break
    return df.rename(columns=mapped)
def dataset_out(d):
    return {'id':d.id,'filename':d.filename,'file_type':d.file_type,'file_size':d.file_size,'records_count':d.records_count,'columns':json.loads(d.columns_json or '[]'),'date_range_start':d.date_range_start,'date_range_end':d.date_range_end,'regions':json.loads(d.regions_json or '[]'),'is_demo':d.is_demo,'created_at':d.created_at}
def metadata(path):
    df=load(path)
    if df.empty: raise ValueError('Uploaded file is empty.')
    dates=pd.to_datetime(df['date'],errors='coerce') if 'date' in df else pd.Series(dtype='datetime64[ns]')
    regions=sorted(df['region'].dropna().astype(str).unique().tolist()) if 'region' in df else []
    return df, {'records_count':len(df),'columns':list(df.columns),'date_range_start': dates.min().strftime('%Y-%m-%d') if not dates.empty and dates.notna().any() else None,'date_range_end':dates.max().strftime('%Y-%m-%d') if not dates.empty and dates.notna().any() else None,'regions':regions}'''

new_load_meta = '''def load(path, nrows=None):
    if path.lower().endswith(('.xlsx','.xls')):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = wb.active
            rows_iter = sheet.iter_rows(values_only=True)
            header_row = next(rows_iter)
            cols = [str(c).strip() for c in header_row if c is not None]
            samples = []
            for i, r in enumerate(rows_iter):
                samples.append(r)
                if nrows is not None and i >= nrows: break
            df = pd.DataFrame(samples, columns=cols[:len(samples[0])] if samples else cols)
        except Exception:
            df = pd.read_excel(path, nrows=nrows)
    else:
        df = pd.read_csv(path, nrows=nrows)
    mapped = {}
    seen_targets = set()
    for col in df.columns:
        n = norm(col)
        for target_name, names in ALIASES.items():
            if (n in names or any(x in n for x in names)) and target_name not in seen_targets:
                mapped[col] = target_name
                seen_targets.add(target_name)
                break
    return df.rename(columns=mapped)

def dataset_out(d):
    return {'id':d.id,'filename':d.filename,'file_type':d.file_type,'file_size':d.file_size,'records_count':d.records_count,'columns':json.loads(d.columns_json or '[]'),'date_range_start':d.date_range_start,'date_range_end':d.date_range_end,'regions':json.loads(d.regions_json or '[]'),'is_demo':d.is_demo,'created_at':d.created_at}

def metadata(path):
    if path.lower().endswith(('.xlsx','.xls')):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            records_count = max(0, (wb.active.max_row or 1) - 1)
        except Exception:
            records_count = len(pd.read_excel(path, usecols=[0]))
    else:
        with open(path, 'rb') as f:
            records_count = max(0, sum(1 for _ in f) - 1)

    df = load(path, nrows=1000)
    if df.empty: raise ValueError('Uploaded file is empty.')
    dates = pd.to_datetime(df['date'], errors='coerce') if 'date' in df else pd.Series(dtype='datetime64[ns]')
    regions = sorted(df['region'].dropna().astype(str).unique().tolist())[:30] if 'region' in df else []
    return df, {'records_count':records_count,'columns':list(df.columns),'date_range_start': dates.min().strftime('%Y-%m-%d') if not dates.empty and dates.notna().any() else None,'date_range_end':dates.max().strftime('%Y-%m-%d') if not dates.empty and dates.notna().any() else None,'regions':regions}'''

text = text.replace(old_load_meta, new_load_meta)

# 3. Update size limit check from 50MB to 100MB
text = text.replace("50*1024*1024", "100*1024*1024")
text = text.replace("File exceeds the 50 MB limit.", "File exceeds the 100 MB limit.")

with open(target, "w", encoding="utf-8") as f:
    f.write(text)

print("SUCCESS: core.py patched for 100MB limit, fast streaming metadata, and column deduplication!")
